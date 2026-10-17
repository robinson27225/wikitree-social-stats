from datetime import datetime

import openpyxl
import pytest

from wtstats import db
from wtstats.importers import find_imports, import_file, remove_import


@pytest.fixture
def env(tmp_path):
    return db.connect(tmp_path / "t.db"), tmp_path / "raw", tmp_path


def rows(conn, sql, *a):
    return conn.execute(sql, a).fetchall()


def make_workbook(path):
    wb = openpyxl.Workbook()
    fb = wb.active
    fb.title = "Copy of FB 2021Daily"
    fb.append(["Date", "Reach", "Page Visits", "New Likes"])
    fb.append([datetime(2021, 1, 1), 111, 9, 3])      # a real Excel date cell
    fb.append(["1/2/2021", "786", 10, 4])            # a date typed as text, not zero-padded
    fb.append(["1/3/2021", None, None, None])        # pre-filled empty row
    ig = wb.create_sheet("Copy of IG 2023Daily")
    ig.append(["Date", "Reach", "Profile Visits", "New Followers"])
    ig.append(["1/1/2023", 163, 2, 2])
    ig.append(["1/3/2023", 123, 3])                  # trailing blank cell left out of the row
    tt = wb.create_sheet("Copy of TT 2023Daily")
    tt.append(["Date", "Video views", "Profile views", "Likes", "Comments", "Shares", "New Followers"])
    tt.append(["2023-09-12", 21, 2, 0, 1, 0, 3])
    li = wb.create_sheet("LI 2023Daily")
    li.append(["Date", "Impressions (total)", "Clicks (total)", "Reactions (total)", "Comments (total)", "Reposts (total)",
               "Engagement rate (total)", "Overview page views (total)", "Overview unique visitors (total)",
               "Total page views (total)", "Total unique visitors (total)", "followers", None, None])
    li.append(["05/04/2023", 0, 0, 0, 0, 0, 0, 2, 2, 3, 2, 1, None, None, 0])  # stray unlabeled cell at the end
    tw = wb.create_sheet("Copy of WTC TW 2024Daily")
    tw.append(["Date", "Tweets published", "impressions", "engagements", "engagement rate"])
    tw.append(["2024-01-01", 0, 6, 1, "16.67%"])
    wb.create_sheet("Notes").append(["just a scratch tab"])
    wb.save(path)


def test_workbook_loads_channel_tabs_and_leaves_twitter_and_junk_out(env):
    conn, raw, tmp = env
    xlsx = tmp / "Archive.xlsx"
    make_workbook(xlsx)
    r = import_file(conn, xlsx, raw_dir=raw)
    status = {s["sheet"]: s["status"] for s in r["sheets"]}
    assert status == {"Copy of FB 2021Daily": "imported", "Copy of IG 2023Daily": "imported", "Copy of TT 2023Daily": "imported",
                      "LI 2023Daily": "imported", "Copy of WTC TW 2024Daily": "skipped", "Notes": "skipped"}
    assert "Twitter" in next(s["reason"] for s in r["sheets"] if s["sheet"].endswith("TW 2024Daily"))
    assert rows(conn, "SELECT COUNT(*) FROM daily_metrics WHERE platform NOT IN ('facebook','instagram','tiktok','linkedin')") == [(0,)]
    # Excel date cell and text date both land as ISO days; the empty prefilled row is dropped
    assert rows(conn, "SELECT period_start FROM daily_metrics WHERE platform='facebook' AND metric='reach' ORDER BY 1") == [
        ("2021-01-01",), ("2021-01-02",)]
    assert rows(conn, "SELECT value FROM daily_metrics WHERE platform='instagram' AND period_start='2023-01-03' AND metric='new_followers'") == [(0.0,)]


def test_old_header_names_map_onto_the_current_ones(env):
    conn, raw, tmp = env
    make_workbook(tmp / "Archive.xlsx")
    import_file(conn, tmp / "Archive.xlsx", raw_dir=raw)
    li = {m for (m,) in rows(conn, "SELECT DISTINCT metric FROM daily_metrics WHERE platform='linkedin'")}
    assert {"impressions", "overview_page_views", "total_unique_visitors", "followers"} <= li and "engagement_rate" not in li
    assert rows(conn, "SELECT value FROM daily_metrics WHERE platform='tiktok' AND metric='followers_change'") == [(3.0,)]  # was "New Followers"


def test_workbook_is_idempotent_and_archived_once(env):
    conn, raw, tmp = env
    make_workbook(tmp / "Archive.xlsx")
    first = import_file(conn, tmp / "Archive.xlsx", raw_dir=raw)
    again = import_file(conn, tmp / "Archive.xlsx", raw_dir=raw)
    assert again["status"] == "skipped" and all(s["status"] == "skipped" for s in again["sheets"])
    assert len(list(raw.glob("*_Archive.xlsx"))) == 1
    assert len(rows(conn, "SELECT id FROM import_log")) == 4  # one per loaded tab


def test_removing_one_tab_keeps_the_shared_archive_until_the_last(env):
    conn, raw, tmp = env
    make_workbook(tmp / "Archive.xlsx")
    import_file(conn, tmp / "Archive.xlsx", raw_dir=raw)
    for n, name in enumerate(["FB 2021Daily", "IG 2023Daily", "TT 2023Daily", "LI 2023Daily"], start=1):
        (found,) = find_imports(conn, name)
        remove_import(conn, found[0])
        assert bool(list(raw.glob("*_Archive.xlsx"))) == (n < 4)
    assert rows(conn, "SELECT COUNT(*) FROM daily_metrics") == [(0,)]


def test_platform_flag_names_an_unrecognised_daily_csv(env):
    conn, raw, tmp = env
    f = tmp / "odd.csv"
    f.write_text("Date,Visits Only\n01/01/2021,5\n")
    r = import_file(conn, f, raw_dir=raw, platform="facebook")
    assert r["platform"] == "facebook" and rows(conn, "SELECT metric FROM daily_metrics") == [("visits_only",)]


def test_repeated_header_and_date_only_block_below_the_data_are_ignored(env):
    conn, raw, tmp = env
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Copy of TT 2024Daily"
    ws.append(["Date", "Video views", "Profile views", "Likes", "Comments", "Shares", "Followers"])
    ws.append(["2024-01-01", 48, 5, 1, 2, 0, 0])
    ws.append([])
    ws.append(["Date"])                      # a stray second header
    ws.append(["2024-01-01"])                # followed by dates with nothing beside them
    ws.append(["2024-01-02"])
    wb.save(tmp / "tt.xlsx")
    r = import_file(conn, tmp / "tt.xlsx", raw_dir=raw)
    assert r["sheets"][0]["rows_loaded"] == 1
    assert rows(conn, "SELECT COUNT(DISTINCT period_start) FROM daily_metrics") == [(1,)]


def test_archived_workbook_keeps_only_the_tabs_that_were_loaded(env):
    conn, raw, tmp = env
    make_workbook(tmp / "Archive.xlsx")
    import_file(conn, tmp / "Archive.xlsx", raw_dir=raw)
    (archived,) = raw.glob("*_Archive.xlsx")
    kept = openpyxl.load_workbook(archived, read_only=True).sheetnames
    assert kept == ["Copy of FB 2021Daily", "Copy of IG 2023Daily", "Copy of TT 2023Daily", "LI 2023Daily"]  # no Twitter, no scratch tab
    assert len(openpyxl.load_workbook(tmp / "Archive.xlsx", read_only=True).sheetnames) == 6  # the original is untouched
    # the archived copy loads to exactly the same data
    from wtstats import db as _db
    other = _db.connect(tmp / "other.db")
    import_file(other, archived, raw_dir=tmp / "raw2")
    q = "SELECT platform, period_start, metric, value FROM daily_metrics ORDER BY 1, 2, 3"
    assert other.execute(q).fetchall() == conn.execute(q).fetchall()


def test_linkedin_blank_cells_are_zero(env):
    conn, raw, tmp = env
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LI 2025Daily"
    ws.append(["Date", "Impressions ", "Clicks ", "Reactions ", "Comments ", "Reposts ", "Engagement rate ", "Overview page views (total)",
               "Overview unique visitors (total)", "Total page views (total)", "Total unique visitors (total)", "followers"])
    ws.append(["07/30/2025", 20, 1, 2, 1, 0, 0.1, 1, 1, 1, 1, None])
    ws.append(["07/31/2025", 18, 0, 1, None, None, 0.05, 0, 0, 0, 0, None])   # comments and reposts blank; followers never filled in
    wb.save(tmp / "li.xlsx")
    import_file(conn, tmp / "li.xlsx", raw_dir=raw)
    assert rows(conn, "SELECT metric, value FROM daily_metrics WHERE platform='linkedin' AND period_start='2025-07-31' "
                      "AND metric IN ('comments','reposts') ORDER BY 1") == [("comments", 0.0), ("reposts", 0.0)]
    # the manually entered followers column is not zero-filled: blank there means "not recorded"
    assert rows(conn, "SELECT COUNT(*) FROM daily_metrics WHERE platform='linkedin' AND metric='followers'") == [(0,)]


def test_linkedin_follower_export_xls_loads_as_followers_and_skips_the_demographic_tabs(env):
    import shutil
    conn, raw, tmp = env
    shutil.copy("tests/fixtures/linkedin_followers_synthetic.xls", tmp / "wikitree_followers_123.xls")
    r = import_file(conn, tmp / "wikitree_followers_123.xls", raw_dir=raw)
    assert {s["sheet"]: s["status"] for s in r["sheets"]} == {"New followers": "imported", "Location": "skipped"}
    assert rows(conn, "SELECT period_start, metric, value FROM daily_metrics ORDER BY 1") == [
        ("2025-10-16", "followers", 0.0), ("2025-10-17", "followers", 2.0), ("2025-10-18", "followers", -1.0), ("2025-10-19", "followers", 3.0)]
    (archived,) = raw.glob("*_wikitree_followers_123.xlsx")           # an .xls is archived as .xlsx, loaded tab only
    assert openpyxl.load_workbook(archived, read_only=True).sheetnames == ["New followers"]
